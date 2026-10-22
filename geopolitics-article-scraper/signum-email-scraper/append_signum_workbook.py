"""Append stage for Signum Iran emails: reads the exported JSON, parses each email, and
appends new rows to a PRIVATE workbook — deliberately NOT the shared PP+G team workbook,
since Signum's reports are your own paid/licensed subscription, not public web content like
the think-tank sources. (Revisit this if you decide the subscription terms allow sharing it.)

Reuses the same repository schema, dedup, sort, and row-formatting logic as the web-scraper's
append_article_repository_workbook.py (imported directly, not duplicated) so both tools stay
consistent and any future fix to one benefits the other automatically.
"""

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import append_article_repository_workbook as appender
import signum_parser


def parse_args():
    ap = argparse.ArgumentParser(description="Append parsed Signum Iran emails to a private repository workbook.")
    ap.add_argument("--input-json", required=True)
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--runner", default="")
    return ap.parse_args()


def emails_to_articles(emails: list) -> list:
    articles = []
    for email in emails:
        parsed = signum_parser.parse_signum_email(email["subject"], email["body"], email["received_date"])
        full_text = parsed["full_text"]
        articles.append({
            "source": "Signum",
            "title": parsed["title"],
            "url": "",  # no public URL for an email; dedup falls back to source+title+date
            "published": parsed["published"],
            "summary": full_text[:300],
            "full_text": full_text,
            "geography": parsed["geography"],
            "geography_code": parsed["geography_code"],
            "topic": parsed["topic"],
        })
    return articles


def main():
    args = parse_args()
    input_json = Path(args.input_json)
    workbook_path = Path(args.workbook)

    raw = appender.retry_on_lock(lambda: input_json.read_text(encoding="utf-8-sig"), "read export JSON").strip()
    emails = json.loads(raw) if raw else []
    if isinstance(emails, dict):
        emails = [emails]
    articles = emails_to_articles(emails)

    workbook = appender.get_or_create_workbook(workbook_path)
    repo_sheet = appender.get_or_create_sheet(workbook, "Repository Rows")
    run_log_sheet = appender.get_or_create_sheet(workbook, "Run Log")
    appender.ensure_header(repo_sheet, appender.REPOSITORY_COLUMNS)
    appender.ensure_header(run_log_sheet, appender.RUN_LOG_COLUMNS)

    added, skipped = appender.append_rows(repo_sheet, articles)
    appender.sort_by_date_desc(repo_sheet)

    run_timestamp = dt.datetime.now().isoformat(timespec="seconds")
    run_log_sheet.append([
        run_timestamp, args.runner, "Iran", str(input_json),
        len(articles), len(added), skipped,
        json.dumps({"Signum": len(added)}),
    ])

    appender.autosize(repo_sheet)
    date_i = appender.REPOSITORY_COLUMNS.index("Publication Date")
    text_i = appender.REPOSITORY_COLUMNS.index("Full Text")
    from openpyxl.styles import Alignment
    for row_num in range(2, repo_sheet.max_row + 1):
        row = repo_sheet[row_num]
        row[date_i].number_format = "m/d/yyyy"
        row[text_i].alignment = Alignment(wrap_text=True, vertical="top")
        repo_sheet.row_dimensions[row_num].height = appender.DATA_ROW_HEIGHT

    appender.retry_on_lock(lambda: workbook.save(workbook_path), "save workbook")

    print(json.dumps({
        "workbook": str(workbook_path),
        "emails_in_json": len(emails),
        "rows_added": len(added),
        "rows_skipped_duplicate": skipped,
    }, indent=2))


if __name__ == "__main__":
    main()
